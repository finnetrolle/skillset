#!/usr/bin/env python3
"""Extract readable text from common CV and vacancy formats."""

from __future__ import annotations

import argparse
import json
from pathlib import Path


def extract_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as exc:
        raise RuntimeError("PDF extraction requires pdfplumber") from exc

    with pdfplumber.open(path) as pdf:
        pages = [(page.extract_text() or "").strip() for page in pdf.pages]
    return "\n\n".join(page for page in pages if page)


def extract_docx(path: Path) -> str:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise RuntimeError("DOCX extraction requires python-docx") from exc

    document = Document(path)
    content = (
        document.iter_inner_content()
        if hasattr(document, "iter_inner_content")
        else [*document.paragraphs, *document.tables]
    )
    blocks: list[str] = []
    for block in content:
        if isinstance(block, Paragraph):
            if block.text.strip():
                blocks.append(block.text.strip())
        elif isinstance(block, Table):
            for row in block.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
    return "\n".join(block for block in blocks if block)


def extract_rtf(path: Path) -> str:
    try:
        from striprtf.striprtf import rtf_to_text
    except ImportError as exc:
        raise RuntimeError("RTF extraction requires striprtf") from exc

    return rtf_to_text(path.read_text(encoding="utf-8", errors="replace")).strip()


def extract(path: Path) -> str:
    suffix = path.suffix.casefold()
    if suffix == ".pdf":
        return extract_pdf(path)
    if suffix == ".docx":
        return extract_docx(path)
    if suffix == ".rtf":
        return extract_rtf(path)
    if suffix == ".json":
        value = json.loads(path.read_text(encoding="utf-8"))
        return json.dumps(value, ensure_ascii=False, indent=2)
    if suffix in {".txt", ".md"}:
        return path.read_text(encoding="utf-8")
    raise ValueError(f"Unsupported input format: {suffix or '<none>'}")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()

    if not args.input.is_file():
        parser.error(f"Input file does not exist: {args.input}")

    try:
        text = extract(args.input)
    except (OSError, RuntimeError, ValueError, json.JSONDecodeError) as exc:
        raise SystemExit(f"Extraction failed: {exc}") from exc
    if not text.strip():
        raise SystemExit("No text extracted; the document may require OCR")

    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(text, encoding="utf-8")
    else:
        print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
