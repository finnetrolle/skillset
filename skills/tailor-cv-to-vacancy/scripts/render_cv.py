#!/usr/bin/env python3
"""Render canonical targeted-CV JSON as a simple ATS-friendly DOCX."""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from typing import Any

from validate_cv import load_and_validate

EN_LABELS = {
    "summary": "Professional summary",
    "skills": "Skills",
    "experience": "Experience",
    "project": "Project",
    "client": "Client",
    "stack": "Stack",
    "responsibilities": "Responsibilities",
    "achievements": "Achievements",
    "education": "Education",
    "courses": "Courses",
    "languages": "Languages",
}
RU_LABELS = {
    "summary": "Профессиональный профиль",
    "skills": "Навыки",
    "experience": "Опыт работы",
    "project": "Проект",
    "client": "Клиент",
    "stack": "Стек",
    "responsibilities": "Обязанности",
    "achievements": "Достижения",
    "education": "Образование",
    "courses": "Курсы",
    "languages": "Языки",
}


def clean_item(value: str) -> str:
    return re.sub(r"^\s*[-•·]\s*", "", value).strip()


def add_lines(document: Any, value: str | None) -> None:
    if not value:
        return
    raw_lines = re.split(r"(?:\r?\n)+|(?=[•·])", value)
    lines = [clean_item(line) for line in raw_lines if clean_item(line)]
    for line in lines:
        document.add_paragraph(line, style="List Bullet")


def add_heading(document: Any, text: str) -> None:
    document.add_paragraph(text.upper(), style="Heading 1")


def select_labels(data: dict[str, Any]) -> dict[str, str]:
    text = " ".join(
        str(data.get(field) or "") for field in ("name", "role", "about", "skills")
    )
    cyrillic = len(re.findall(r"[А-Яа-яЁё]", text))
    latin = len(re.findall(r"[A-Za-z]", text))
    return RU_LABELS if cyrillic > latin else EN_LABELS


def render(data: dict[str, Any], output: Path) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("DOCX rendering requires python-docx") from exc

    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    heading_tokens = {
        "Heading 1": (16, "2E74B5", 18, 10),
        "Heading 2": (13, "2E74B5", 14, 7),
        "Heading 3": (12, "1F4D78", 10, 5),
    }
    for style_name, (size, color, before, after) in heading_tokens.items():
        style = document.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    bullet = document.styles["List Bullet"]
    bullet.font.name = "Calibri"
    bullet._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    bullet._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    bullet.font.size = Pt(11)
    bullet.paragraph_format.left_indent = Inches(0.375)
    bullet.paragraph_format.first_line_indent = Inches(-0.188)
    bullet.paragraph_format.space_after = Pt(4)
    bullet.paragraph_format.line_spacing = 1.25

    labels = select_labels(data)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name_run = title.add_run(data.get("name") or "CV")
    name_run.bold = True
    name_run.font.name = "Calibri"
    name_run.font.size = Pt(22)
    title.paragraph_format.space_after = Pt(3)

    role = document.add_paragraph()
    role.alignment = WD_ALIGN_PARAGRAPH.CENTER
    role_run = role.add_run(data.get("role") or "")
    role_run.bold = True
    role_run.font.name = "Calibri"
    role_run.font.size = Pt(12)
    role.paragraph_format.space_after = Pt(3)
    if data.get("location"):
        location = document.add_paragraph(data["location"])
        location.alignment = WD_ALIGN_PARAGRAPH.CENTER
        location.paragraph_format.space_after = Pt(10)

    if data.get("about"):
        add_heading(document, labels["summary"])
        document.add_paragraph(data["about"])
    if data.get("skills"):
        add_heading(document, labels["skills"])
        document.add_paragraph(data["skills"])

    if data["experience"]:
        add_heading(document, labels["experience"])
        for job in data["experience"]:
            heading = document.add_paragraph()
            heading.paragraph_format.keep_with_next = True
            heading.paragraph_format.space_before = Pt(8)
            main = " | ".join(
                value for value in (job.get("position"), job.get("company")) if value
            )
            main_run = heading.add_run(main or "Experience")
            main_run.bold = True
            if job.get("time"):
                heading.add_run(f"\n{job['time']}").italic = True
            if job.get("project"):
                project = document.add_paragraph()
                project.add_run(f"{labels['project']}: ").bold = True
                project.add_run(job["project"])
            if job.get("client"):
                document.add_paragraph(f"{labels['client']}: {job['client']}")
            if job.get("stack"):
                stack = document.add_paragraph()
                stack.add_run(f"{labels['stack']}: ").bold = True
                stack.add_run(job["stack"])
            if job.get("responsibilities"):
                label = document.add_paragraph()
                label.paragraph_format.keep_with_next = True
                label.add_run(labels["responsibilities"]).bold = True
                add_lines(document, job["responsibilities"])
            if job.get("achievements"):
                label = document.add_paragraph()
                label.paragraph_format.keep_with_next = True
                label.add_run(labels["achievements"]).bold = True
                add_lines(document, job["achievements"])

    if data["education"]:
        add_heading(document, labels["education"])
        for item in data["education"]:
            parts = [item.get("name"), item.get("direction")]
            text = " — ".join(str(value) for value in parts if value)
            if item.get("grad_year"):
                text = (
                    f"{text}, {item['grad_year']}" if text else str(item["grad_year"])
                )
            document.add_paragraph(text)

    if data["courses"]:
        add_heading(document, labels["courses"])
        for item in data["courses"]:
            text = item.get("name") or ""
            if item.get("grad_year"):
                text = (
                    f"{text}, {item['grad_year']}" if text else str(item["grad_year"])
                )
            document.add_paragraph(text, style="List Bullet")

    if data["languages"]:
        add_heading(document, labels["languages"])
        for item in data["languages"]:
            text = item.get("name") or ""
            if item.get("level"):
                text = f"{text} - {item['level']}" if text else item["level"]
            document.add_paragraph(text, style="List Bullet")

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    output = args.output or args.input.with_suffix(".docx")
    try:
        data = load_and_validate(args.input)
        render(data, output)
    except (RuntimeError, ValueError) as exc:
        print(exc)
        return 1
    print(f"Created: {output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
